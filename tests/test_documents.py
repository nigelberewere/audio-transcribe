import io
from pathlib import Path

import pytest
from fastapi import HTTPException, UploadFile
from PIL import Image
from docx import Document
import pypdf
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app import main
from app.db import Database
from app.security import hash_password


@pytest.fixture
def document_state(tmp_path, monkeypatch):
    database = Database(tmp_path / "jobs.sqlite3")
    database.create_user("owner", hash_password("owner-pass"), "admin", None)
    database.create_user("member", hash_password("member-pass"), "user", "owner")
    monkeypatch.setattr(main, "database", database)
    monkeypatch.setattr(main.settings, "documents_storage_path", tmp_path / "documents")
    main.settings.documents_storage_path.mkdir(parents=True, exist_ok=True)
    main.sessions.clear()
    return database


def upload(name, content, content_type="text/plain"):
    return UploadFile(io.BytesIO(content), filename=name, headers={"content-type": content_type})


def make_pdf(pages=1, text="Sample"):
    packet = io.BytesIO()
    can = canvas.Canvas(packet, pagesize=letter)
    for i in range(pages):
        can.drawString(100, 700, f"{text} Page {i + 1}")
        can.showPage()
    can.save()
    packet.seek(0)
    return packet.getvalue()


def make_docx(text="Hello docx world"):
    doc = Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_png():
    img = Image.new("RGB", (60, 60), color="blue")
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def test_upload_creates_document_version_and_audit(document_state):
    document = main.upload_document(upload("policy.txt", b"one"), None, "policy, public", "member")
    assert document["current_version"] == 1
    assert document["tags"] == ["policy", "public"]
    assert len(document_state.list_document_versions(document["id"])) == 1
    with document_state.connect() as connection:
        audit = connection.execute("SELECT action, target FROM audit_log").fetchone()
    assert tuple(audit) == ("document_uploaded", document["id"])


def test_new_version_preserves_old_file(document_state):
    document = main.upload_document(upload("minutes.txt", b"old"), None, "", "member")
    old_path = document["storage_path"]
    updated = main.upload_document_version(document["id"], upload("minutes.txt", b"new"), "updated", "member")
    assert updated["current_version"] == 2
    assert open(old_path, "rb").read() == b"old"
    assert len(document_state.list_document_versions(document["id"])) == 2


def test_soft_delete_hides_from_users_and_admin_can_restore(document_state):
    document = main.upload_document(upload("contract.txt", b"contract"), None, "", "member")
    main.delete_document(document["id"], "member")
    assert document_state.list_documents() == []
    assert document_state.list_deleted_documents()[0]["id"] == document["id"]
    restored = main.restore_document(document["id"], "owner")
    assert restored["deleted"] is False
    assert document_state.list_documents()[0]["id"] == document["id"]


def test_document_routes_require_login(document_state):
    with pytest.raises(HTTPException) as error:
        main.current_user(None)
    assert error.value.status_code == 401


# --- Integration Tests: Download, Folders, Tags, and Version History ---

def test_download_round_trip(document_state):
    known_bytes = b"Exact byte sequence with binary \x00\x01\xfe\xff and text: Round-trip verification."
    uploaded = main.upload_document(upload("roundtrip.txt", known_bytes), None, "", "member")
    assert uploaded["id"] is not None

    download_response = main.download_document(uploaded["id"], user="member")
    assert download_response.status_code == 200
    assert download_response.filename == "roundtrip.txt"
    downloaded_bytes = Path(download_response.path).read_bytes()
    assert downloaded_bytes == known_bytes
    assert len(downloaded_bytes) == len(known_bytes)


def test_nested_folder_navigation(document_state):
    parent_folder = main.create_document_folder(name="Projects", parent_folder_id=None, user="member")
    assert parent_folder["id"] is not None
    assert parent_folder["name"] == "Projects"
    assert parent_folder["parent_folder_id"] is None

    nested_folder = main.create_document_folder(name="2026", parent_folder_id=parent_folder["id"], user="member")
    assert nested_folder["id"] is not None
    assert nested_folder["name"] == "2026"
    assert nested_folder["parent_folder_id"] == parent_folder["id"]

    doc = main.upload_document(upload("spec.txt", b"Nested specifications"), folder_id=nested_folder["id"], tags="", user="member")
    assert doc["folder_id"] == nested_folder["id"]

    # Document appears when listing nested folder's contents
    nested_listing = main.list_documents(folder_id=nested_folder["id"], user="member")
    assert any(d["id"] == doc["id"] for d in nested_listing["documents"])
    matched = next(d for d in nested_listing["documents"] if d["id"] == doc["id"])
    assert matched["filename"] == "spec.txt"

    # Document does NOT appear when listing top-level/root folder
    root_listing = main.list_documents(folder_id=None, user="member")
    assert not any(d["id"] == doc["id"] for d in root_listing["documents"])

    # Folder listing endpoint correctly reports nested folder's parent_folder_id for breadcrumb navigation
    parent_listing = main.list_documents(folder_id=parent_folder["id"], user="member")
    child_info = next((f for f in parent_listing["folders"] if f["id"] == nested_folder["id"]), None)
    assert child_info is not None
    assert child_info["parent_folder_id"] == parent_folder["id"]

    # Top-level folder listing reports parent_folder_id as None
    root_parent_info = next((f for f in root_listing["folders"] if f["id"] == parent_folder["id"]), None)
    assert root_parent_info is not None
    assert root_parent_info["parent_folder_id"] is None


def test_tag_add_and_remove_persists(document_state):
    doc = main.upload_document(upload("memo.txt", b"Memo content"), None, "", "member")
    doc_id = doc["id"]
    assert doc["tags"] == []

    # Add tags to the document via API
    res_add = main.update_document_tags(doc_id, tags="urgent, review", user="member")
    assert sorted(res_add["tags"]) == ["review", "urgent"]

    # Fetch document again as if it were a fresh request: tag is present
    fresh_doc = main.get_document(doc_id, user="member")
    assert sorted(fresh_doc["tags"]) == ["review", "urgent"]

    # Folder listing also reflects the persisted tags
    fresh_listing = main.list_documents(folder_id=None, user="member")
    listed_doc = next(d for d in fresh_listing["documents"] if d["id"] == doc_id)
    assert sorted(listed_doc["tags"]) == ["review", "urgent"]

    # Confirm tags are in actual document_tags table
    with document_state.connect() as conn:
        db_tags = [r[0] for r in conn.execute("SELECT tag FROM document_tags WHERE document_id = ? ORDER BY tag", (doc_id,)).fetchall()]
    assert db_tags == ["review", "urgent"]

    # Remove the tags via API
    res_remove = main.update_document_tags(doc_id, tags="", user="member")
    assert res_remove["tags"] == []

    # Subsequent fetch confirms tag is no longer present
    subsequent_doc = main.get_document(doc_id, user="member")
    assert subsequent_doc["tags"] == []

    subsequent_listing = main.list_documents(folder_id=None, user="member")
    subsequent_listed_doc = next(d for d in subsequent_listing["documents"] if d["id"] == doc_id)
    assert subsequent_listed_doc["tags"] == []

    # Confirm document_tags table no longer has rows for this document
    with document_state.connect() as conn:
        remaining_db_tags = conn.execute("SELECT tag FROM document_tags WHERE document_id = ?", (doc_id,)).fetchall()
    assert remaining_db_tags == []


def test_change_note_recorded_and_old_version_downloadable(document_state):
    v1_content = b"Original contract draft version 1"
    v2_content = b"Revised contract draft version 2 with updated liability section"
    note = "Updated liability cap and payment schedule"

    doc = main.upload_document(upload("agreement.txt", v1_content), None, "", "member")
    doc_id = doc["id"]

    v2_doc = main.upload_document_version(doc_id, upload("agreement.txt", v2_content), change_note=note, user="member")
    assert v2_doc["current_version"] == 2

    # Change note is retrievable via version history endpoint
    doc_details = main.get_document(doc_id, user="member")
    assert "versions" in doc_details
    versions = doc_details["versions"]
    assert len(versions) == 2
    v2_record = next(v for v in versions if v["version_number"] == 2)
    assert v2_record["change_note"] == note
    v1_record = next(v for v in versions if v["version_number"] == 1)
    assert v1_record["change_note"] is None

    # OLD version can still be downloaded by its specific version id/number with original content intact
    download_old = main.download_document(doc_id, version=1, user="member")
    assert download_old.status_code == 200
    downloaded_old_bytes = Path(download_old.path).read_bytes()
    assert downloaded_old_bytes == v1_content
    assert downloaded_old_bytes != v2_content

    # Current version download returns new content
    download_current = main.download_document(doc_id, user="member")
    assert download_current.status_code == 200
    assert Path(download_current.path).read_bytes() == v2_content

    # Specific version 2 download returns new content
    download_v2 = main.download_document(doc_id, version=2, user="member")
    assert download_v2.status_code == 200
    assert Path(download_v2.path).read_bytes() == v2_content


# --- Restrictions & Size Limit Tests ---

def test_upload_allowed_type_under_50mb_succeeds(document_state):
    pdf_bytes = make_pdf(pages=1, text="Test Document")
    doc = main.upload_document(upload("report.pdf", pdf_bytes, "application/pdf"), None, "", "member")
    assert doc is not None
    assert doc["filename"] == "report.pdf"
    assert Path(doc["storage_path"]).is_file()
    assert len(document_state.list_documents()) == 1


def test_upload_disallowed_type_rejected_and_leaves_no_traces(document_state):
    storage_root = main.settings.documents_storage_path
    initial_files = list(storage_root.glob("**/*"))

    with pytest.raises(HTTPException) as error:
        main.upload_document(upload("malicious.exe", b"binarycontent", "application/x-msdownload"), None, "", "member")
    assert error.value.status_code == 400
    assert "not allowed" in error.value.detail.lower()

    # Confirm no DB row created
    assert document_state.list_documents() == []
    # Confirm no files written to storage
    current_files = list(storage_root.glob("**/*"))
    assert current_files == initial_files


def test_upload_over_50mb_rejected_and_leaves_no_traces(document_state, monkeypatch):
    storage_root = main.settings.documents_storage_path
    initial_files = list(storage_root.glob("**/*"))

    # Create a mock stream that yields slightly more than 50MB
    class LargeStream(io.BytesIO):
        def __init__(self, size):
            self.total = size
            self.sent = 0

        def read(self, chunk_size=1024 * 1024):
            if self.sent >= self.total:
                return b""
            chunk = min(chunk_size, self.total - self.sent)
            self.sent += chunk
            return b"X" * chunk

    over_50mb = LargeStream(50 * 1024 * 1024 + 1024)
    file_obj = UploadFile(over_50mb, filename="large_report.pdf", headers={"content-type": "application/pdf"})

    with pytest.raises(HTTPException) as error:
        main.upload_document(file_obj, None, "", "member")
    assert error.value.status_code == 400
    assert "50mb" in error.value.detail.lower()

    # Confirm no DB row created
    assert document_state.list_documents() == []
    # Confirm no files written to storage
    current_files = list(storage_root.glob("**/*"))
    assert current_files == initial_files


# --- PDF Tools: Merge Tests ---

def test_merge_two_pdfs_produces_combined_page_count_and_lineage(document_state):
    pdf1 = make_pdf(pages=2, text="Doc One")
    pdf2 = make_pdf(pages=3, text="Doc Two")
    doc1 = main.upload_document(upload("doc1.pdf", pdf1, "application/pdf"), None, "", "member")
    doc2 = main.upload_document(upload("doc2.pdf", pdf2, "application/pdf"), None, "", "member")

    req = main.MergeDocumentsRequest(document_ids=[doc1["id"], doc2["id"]], output_filename="combined.pdf")
    merged = main.merge_documents_endpoint(req, user="member")

    assert merged["filename"] == "combined.pdf"
    assert merged["source_document_ids"] == [doc1["id"], doc2["id"]]
    reader = pypdf.PdfReader(merged["storage_path"])
    assert len(reader.pages) == 5

    # Check audit log
    with document_state.connect() as conn:
        audit = conn.execute("SELECT action, target, details FROM audit_log WHERE action = 'document_merged'").fetchone()
    assert audit is not None
    assert audit[0] == "document_merged"
    assert audit[1] == merged["id"]
    assert doc1["id"] in audit[2] and doc2["id"] in audit[2]


def test_merge_fewer_than_two_files_rejected(document_state):
    pdf = make_pdf(pages=1, text="Solo")
    doc = main.upload_document(upload("solo.pdf", pdf, "application/pdf"), None, "", "member")

    req = main.MergeDocumentsRequest(document_ids=[doc["id"]])
    with pytest.raises(HTTPException) as error:
        main.merge_documents_endpoint(req, user="member")
    assert error.value.status_code == 400
    assert "at least 2" in error.value.detail.lower()


# --- PDF Tools: Split Tests ---

def test_split_pdf_by_page_range(document_state):
    pdf = make_pdf(pages=5, text="MultiPage")
    doc = main.upload_document(upload("multipage.pdf", pdf, "application/pdf"), None, "", "member")

    req = main.SplitDocumentRequest(document_id=doc["id"], page_ranges="1-2, 3-5")
    res = main.split_document_endpoint(req, user="member")

    created = res["documents"]
    assert len(created) == 2
    r1 = pypdf.PdfReader(created[0]["storage_path"])
    assert len(r1.pages) == 2
    r2 = pypdf.PdfReader(created[1]["storage_path"])
    assert len(r2.pages) == 3

    assert created[0]["source_document_ids"] == [doc["id"]]
    assert created[1]["source_document_ids"] == [doc["id"]]

    with document_state.connect() as conn:
        audits = conn.execute("SELECT action, target FROM audit_log WHERE action = 'document_split'").fetchall()
    assert len(audits) == 2


def test_split_more_than_one_file_rejected(document_state):
    pdf1 = make_pdf(pages=1, text="One")
    pdf2 = make_pdf(pages=1, text="Two")
    doc1 = main.upload_document(upload("one.pdf", pdf1, "application/pdf"), None, "", "member")
    doc2 = main.upload_document(upload("two.pdf", pdf2, "application/pdf"), None, "", "member")

    req = main.SplitDocumentRequest(document_ids=[doc1["id"], doc2["id"]])
    with pytest.raises(HTTPException) as error:
        main.split_document_endpoint(req, user="member")
    assert error.value.status_code == 400
    assert "exactly 1" in error.value.detail.lower()


# --- PDF Tools: Watermark Tests ---

def test_watermark_preserves_page_count_and_opens(document_state):
    pdf = make_pdf(pages=3, text="WatermarkSource")
    doc = main.upload_document(upload("source.pdf", pdf, "application/pdf"), None, "", "member")

    req = main.WatermarkDocumentRequest(document_ids=[doc["id"]], watermark_text="CONFIDENTIAL")
    res = main.watermark_documents_endpoint(req, user="member")

    watermarked_doc = res["documents"][0]
    assert watermarked_doc["source_document_ids"] == [doc["id"]]

    reader = pypdf.PdfReader(watermarked_doc["storage_path"])
    assert len(reader.pages) == 3

    with document_state.connect() as conn:
        audit = conn.execute("SELECT action, target, details FROM audit_log WHERE action = 'document_watermarked'").fetchone()
    assert audit[0] == "document_watermarked"
    assert audit[1] == watermarked_doc["id"]
    assert "CONFIDENTIAL" in audit[2]


# --- File Conversion Tests ---

def test_conversion_docx_to_pdf_and_pdf_to_docx(document_state):
    # docx -> pdf
    docx_bytes = make_docx("Sample text inside Word document")
    doc_word = main.upload_document(upload("report.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), None, "", "member")

    req_pdf = main.ConvertDocumentRequest(document_id=doc_word["id"], target_format="pdf")
    converted_pdf = main.convert_document_endpoint(req_pdf, user="member")
    assert converted_pdf["filename"] == "report.pdf"
    assert converted_pdf["source_document_ids"] == [doc_word["id"]]
    reader = pypdf.PdfReader(converted_pdf["storage_path"])
    assert len(reader.pages) >= 1

    # pdf -> docx
    req_docx = main.ConvertDocumentRequest(document_id=converted_pdf["id"], target_format="docx")
    converted_docx = main.convert_document_endpoint(req_docx, user="member")
    assert converted_docx["filename"] == "report.docx"
    assert converted_docx["source_document_ids"] == [converted_pdf["id"]]
    doc_check = Document(converted_docx["storage_path"])
    assert any("Sample text" in p.text for p in doc_check.paragraphs)


def test_conversion_image_to_pdf(document_state):
    png_bytes = make_png()
    doc_img = main.upload_document(upload("scan.png", png_bytes, "image/png"), None, "", "member")

    req = main.ConvertDocumentRequest(document_id=doc_img["id"], target_format="pdf")
    converted = main.convert_document_endpoint(req, user="member")
    assert converted["filename"] == "scan.pdf"
    assert converted["source_document_ids"] == [doc_img["id"]]
    reader = pypdf.PdfReader(converted["storage_path"])
    assert len(reader.pages) == 1

    with document_state.connect() as conn:
        audit = conn.execute("SELECT action, target, details FROM audit_log WHERE action = 'document_converted' AND target = ?", (converted["id"],)).fetchone()
    assert audit is not None
    assert "pdf" in audit[2]

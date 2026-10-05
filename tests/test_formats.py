from app.formats import format_timestamp, group_segments_into_paragraphs, timestamp, write_outputs


def test_timestamp_is_hour_safe():
    assert timestamp(0) == "00:00:00"
    assert timestamp(3661.9) == "01:01:01"


def test_timestamp_preserves_milliseconds_for_subtitle_formats():
    assert format_timestamp(83.421, "srt") == "00:01:23,421"
    assert format_timestamp(83.421, "vtt") == "00:01:23.421"


def test_paragraphs_merge_until_pause():
    segments = [
        {"start": 0.0, "end": 1.0, "text": "First."},
        {"start": 1.5, "end": 2.0, "text": "Second."},
        {"start": 5.0, "end": 6.0, "text": "Third."},
    ]
    paragraphs = group_segments_into_paragraphs(segments)
    assert [paragraph["text"] for paragraph in paragraphs] == ["First. Second.", "Third."]


def test_paragraphs_split_when_speaker_changes():
    segments = [
        {"start": 0.0, "end": 1.0, "text": "One.", "speaker": "Speaker 1"},
        {"start": 1.2, "end": 2.0, "text": "Two.", "speaker": "Speaker 2"},
    ]
    paragraphs = group_segments_into_paragraphs(segments)
    assert len(paragraphs) == 2


def test_docx_contains_metadata_table(tmp_path):
    job = {
        "id": "job-1",
        "filename": "meeting.wav",
        "formats": ["docx"],
        "selected_model": "medium",
        "duration": 12.5,
        "device": "cpu",
        "compute_type": "int8",
        "detected_language": "en",
        "completed_at": "2026-09-23T12:00:00",
    }
    write_outputs(job, [{"start": 0.0, "end": 1.0, "text": "Hello."}], tmp_path)

    from docx import Document

    document = Document(tmp_path / "meeting.wav.docx" if (tmp_path / "meeting.wav.docx").exists() else tmp_path / "meeting.docx")
    rows = {(row.cells[0].text, row.cells[1].text) for row in document.tables[0].rows}
    assert ("Filename", "meeting.wav") in rows
    assert ("Date processed", "Sep 23, 2026, 12:00 PM") in rows
    assert ("Total segment count", "1") in rows


def test_docx_metadata_table_never_contains_literal_none(tmp_path):
    job_with_nones = {
        "id": "job-none",
        "filename": "recording.mp3",
        "formats": ["docx"],
        "completed_at": None,
        "processed_at": None,
        "updated_at": None,
        "created_at": None,
        "duration": None,
        "selected_model": None,
        "requested_model": None,
        "device": None,
        "compute_type": None,
        "detected_language": None,
        "language": None,
    }
    write_outputs(job_with_nones, [{"start": 0.0, "end": 2.0, "text": "Testing none values."}], tmp_path)

    from docx import Document

    document = Document(tmp_path / "recording.docx")
    table = document.tables[0]

    row_dict = {}
    for row in table.rows:
        label = row.cells[0].text
        value = row.cells[1].text
        row_dict[label] = value
        assert label != "None"
        assert value != "None"
        assert "None" not in value.split()
        assert label.strip() != ""
        assert value.strip() != ""

    assert "Date processed" in row_dict
    assert row_dict["Date processed"] != "None"
    assert any(m in row_dict["Date processed"] for m in ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"))


def test_speaker_names_are_applied_to_exports_without_changing_segments(tmp_path):
    job = {
        "id": "job-speakers",
        "filename": "meeting.wav",
        "formats": ["txt", "srt", "docx"],
        "selected_model": "medium",
        "duration": 4.0,
        "completed_at": "2026-09-23T12:00:00",
    }
    segments = [
        {"start": 0.0, "end": 1.0, "speaker": "Speaker 1", "text": "Welcome."},
        {"start": 1.2, "end": 2.0, "speaker": "Speaker 2", "text": "Thank you."},
    ]

    write_outputs(job, segments, tmp_path, {"Speaker 1": "Alice"})

    assert "Alice: Welcome." in (tmp_path / "meeting.txt").read_text(encoding="utf-8")
    assert "Speaker 2: Thank you." in (tmp_path / "meeting.txt").read_text(encoding="utf-8")
    assert "Alice: Welcome." in (tmp_path / "meeting.srt").read_text(encoding="utf-8")
    assert "Speaker 2: Thank you." in (tmp_path / "meeting.srt").read_text(encoding="utf-8")
    from docx import Document
    docx_text = "\n".join(paragraph.text for paragraph in Document(tmp_path / "meeting.docx").paragraphs)
    assert "Alice: Welcome." in docx_text
    assert "Speaker 2: Thank you." in docx_text
    assert segments[0]["speaker"] == "Speaker 1"


def test_format_datetime_variations():
    from datetime import datetime
    from app.formats import format_datetime

    # Explicit naive string
    assert format_datetime("2026-09-26T15:42:00") == "Sep 26, 2026, 3:42 PM"
    assert format_datetime("2026-09-26 09:05:00") == "Sep 26, 2026, 9:05 AM"

    # Datetime object
    assert format_datetime(datetime(2026, 9, 26, 15, 42)) == "Sep 26, 2026, 3:42 PM"

    # None and empty string fallback safely without error or literal "None"
    res_none = format_datetime(None)
    assert res_none != "None"
    assert ", " in res_none

    res_empty = format_datetime("")
    assert res_empty != "None"

    res_literal = format_datetime("None")
    assert res_literal != "None"

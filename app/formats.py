import json
from datetime import datetime
from pathlib import Path
from typing import Any


def format_timestamp(seconds: float, fmt: str = "display") -> str:
    total_msec = int(round(seconds * 1000))
    hours = total_msec // 3_600_000
    minutes = (total_msec % 3_600_000) // 60_000
    secs = (total_msec % 60_000) // 1000
    msec = total_msec % 1000
    if fmt == "srt":
        return f"{hours:02d}:{minutes:02d}:{secs:02d},{msec:03d}"
    elif fmt == "vtt":
        return f"{hours:02d}:{minutes:02d}:{secs:02d}.{msec:03d}"
    else:
        return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def timestamp(seconds: float, fmt: str = "display") -> str:
    """Backward-compatible alias for callers using the old formatter name."""
    return format_timestamp(seconds, fmt)


def group_segments_into_paragraphs(
    segments: list[dict], pause_threshold: float = 2.0
) -> list[dict]:
    """Merge consecutive segments until a long pause or speaker change."""
    paragraphs: list[dict] = []
    current: dict[str, Any] | None = None
    has_speakers = _has_speakers(segments)
    for segment in segments:
        text = segment.get("text", "").strip()
        if not text:
            continue
        start = float(segment.get("start", 0.0))
        end = float(segment.get("end", start))
        speaker = segment.get("speaker")
        can_merge = (
            current is not None
            and start - current["end"] <= pause_threshold
            and (not has_speakers or speaker == current["speaker"])
        )
        if can_merge:
            current["texts"].append(text)
            current["end"] = end
        else:
            if current is not None:
                paragraphs.append({**current, "text": " ".join(current.pop("texts")).strip()})
            current = {"start": start, "end": end, "speaker": speaker, "texts": [text]}
    if current is not None:
        paragraphs.append({**current, "text": " ".join(current.pop("texts")).strip()})
    return paragraphs


def _has_speakers(segments: list[dict]) -> bool:
    return any(segment.get("speaker") for segment in segments)


def _display_text(paragraph: dict, has_speakers: bool, speaker_names: dict[str, str] | None = None) -> str:
    speaker = paragraph.get("speaker")
    display_name = speaker_names.get(speaker, speaker) if speaker_names and speaker else speaker
    prefix = f"{display_name}: " if has_speakers and display_name else ""
    return f"{prefix}{paragraph['text']}"


def format_datetime(value: Any = None) -> str:
    """Format an ISO timestamp or datetime into a user-friendly string (e.g., 'Sep 26, 2026, 3:42 PM')."""
    if value is None:
        dt = datetime.now()
    elif isinstance(value, (int, float)):
        dt = datetime.fromtimestamp(value)
    elif isinstance(value, str):
        val = value.strip()
        if not val or val.lower() == "none":
            dt = datetime.now()
        else:
            try:
                dt = datetime.fromisoformat(val.replace("Z", "+00:00"))
            except Exception:
                try:
                    dt = datetime.strptime(val, "%Y-%m-%d %H:%M:%S")
                except Exception:
                    dt = datetime.now()
    elif isinstance(value, datetime):
        dt = value
    else:
        dt = datetime.now()

    if dt.tzinfo is not None:
        dt = dt.astimezone()

    date_str = f"{dt.strftime('%b')} {dt.day}, {dt.year}"
    hour = dt.strftime("%I").lstrip("0") or "12"
    minute_ampm = dt.strftime("%M %p")
    return f"{date_str}, {hour}:{minute_ampm}"


def write_outputs(
    job: dict,
    segments: list[dict],
    output_dir: Path,
    speaker_names: dict[str, str] | None = None,
) -> list[Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(job["filename"]).stem
    paths = []
    paragraphs = group_segments_into_paragraphs(segments)
    has_speakers = _has_speakers(segments)
    plain = "\n\n".join(_display_text(paragraph, has_speakers, speaker_names) for paragraph in paragraphs)
    timed = "\n\n".join(f"[{format_timestamp(paragraph['start'])}] {_display_text(paragraph, has_speakers, speaker_names)}" for paragraph in paragraphs)
    if "txt" in job["formats"]:
        path = output_dir / f"{stem}.txt"; path.write_text(plain + "\n", encoding="utf-8"); paths.append(path)
    if "txt_timestamps" in job["formats"]:
        path = output_dir / f"{stem}_timestamps.txt"; path.write_text(timed + "\n", encoding="utf-8"); paths.append(path)
    if "srt" in job["formats"]:
        body = "\n\n".join(f"{index}\n{format_timestamp(segment.get('start', 0), 'srt')} --> {format_timestamp(segment.get('end', 0), 'srt')}\n{_display_text(segment, has_speakers, speaker_names)}" for index, segment in enumerate(segments, 1) if segment.get("text", "").strip())
        path = output_dir / f"{stem}.srt"; path.write_text(body, encoding="utf-8"); paths.append(path)
    if "vtt" in job["formats"]:
        body = "WEBVTT\n\n" + "\n\n".join(f"{format_timestamp(segment.get('start', 0), 'vtt')} --> {format_timestamp(segment.get('end', 0), 'vtt')}\n{_display_text(segment, has_speakers, speaker_names)}" for segment in segments if segment.get("text", "").strip())
        path = output_dir / f"{stem}.vtt"; path.write_text(body, encoding="utf-8"); paths.append(path)
    if "json" in job["formats"]:
        export_segments = [
            {**segment, "speaker": speaker_names.get(segment.get("speaker"), segment.get("speaker"))}
            if speaker_names and segment.get("speaker") else dict(segment)
            for segment in segments
        ]
        path = output_dir / f"{stem}.json"; path.write_text(json.dumps({"job_id": job["id"], "filename": job["filename"], "model": job["selected_model"], "segments": export_segments}, indent=2), encoding="utf-8"); paths.append(path)
    if "docx" in job["formats"]:
        from docx import Document
        from docx.shared import RGBColor
        document = Document(); document.add_heading(job["filename"], 0)
        raw_date = (
            job.get("completed_at")
            or job.get("processed_at")
            or job.get("updated_at")
            or job.get("created_at")
        )
        date_processed = format_datetime(raw_date)
        metadata = [
            ("Filename", job.get("filename") or "Unknown"),
            ("Date processed", date_processed),
            ("Audio duration", format_timestamp(float(job.get("duration") or 0.0))),
            ("Model used", job.get("selected_model") or job.get("requested_model") or "unknown"),
            ("Device / compute type", f"{job.get('device') or 'cpu'} / {job.get('compute_type') or 'int8'}"),
            ("Detected language", job.get("detected_language") or job.get("language") or "auto"),
            ("Total segment count", str(len(segments))),
        ]
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for label, value in metadata:
            cells = table.add_row().cells
            cells[0].text = label
            val_str = str(value) if value is not None and str(value).strip().lower() != "none" else "—"
            cells[1].text = val_str
        document.add_paragraph()
        for paragraph in paragraphs:
            para = document.add_paragraph()
            timestamp_run = para.add_run(f"[{format_timestamp(paragraph['start'])}] ")
            timestamp_run.font.color.rgb = RGBColor(100, 116, 139)
            if has_speakers and paragraph.get("speaker"):
                speaker = speaker_names.get(paragraph["speaker"], paragraph["speaker"]) if speaker_names else paragraph["speaker"]
                speaker_run = para.add_run(f"{speaker}: ")
                speaker_run.bold = True
                speaker_run.font.color.rgb = RGBColor(30, 58, 138)
            para.add_run(paragraph["text"])
        path = output_dir / f"{stem}.docx"; document.save(path); paths.append(path)
    return paths

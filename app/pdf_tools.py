import io
import json
import xml.sax.saxutils as saxutils
from pathlib import Path
from typing import Any

from PIL import Image
from docx import Document
import pypdf
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def merge_pdfs(input_paths: list[Path], output_path: Path) -> int:
    """Merge multiple PDF files into a single PDF. Returns total page count."""
    if len(input_paths) < 2:
        raise ValueError("Merge requires at least 2 PDF documents.")
    writer = pypdf.PdfWriter()
    for path in input_paths:
        reader = pypdf.PdfReader(str(path))
        for page in reader.pages:
            writer.add_page(page)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("wb") as f:
        writer.write(f)
    return len(writer.pages)


def parse_page_ranges(ranges_str: str | None, total_pages: int) -> list[list[int]]:
    """
    Parse page ranges into 0-based page index lists.
    Examples:
      - None, '', 'all', 'each' -> [[0], [1], ..., [total_pages-1]]
      - '1-3, 5' -> [[0, 1, 2], [4]]
    """
    if not ranges_str or ranges_str.strip().lower() in ("all", "each"):
        return [[i] for i in range(total_pages)]

    parts = [p.strip() for p in ranges_str.split(",") if p.strip()]
    if not parts:
        return [[i] for i in range(total_pages)]

    result: list[list[int]] = []
    for part in parts:
        if "-" in part:
            start_str, end_str = part.split("-", 1)
            try:
                start = int(start_str.strip())
                end = int(end_str.strip())
            except ValueError:
                raise ValueError(f"Invalid page range format: '{part}'.")
            if start < 1 or end < start or end > total_pages:
                raise ValueError(f"Page range '{part}' is out of bounds (document has {total_pages} pages).")
            result.append(list(range(start - 1, end)))
        else:
            try:
                page_num = int(part)
            except ValueError:
                raise ValueError(f"Invalid page number: '{part}'.")
            if page_num < 1 or page_num > total_pages:
                raise ValueError(f"Page number {page_num} is out of bounds (document has {total_pages} pages).")
            result.append([page_num - 1])
    return result


def split_pdf(
    input_path: Path, output_dir: Path, base_filename: str, page_ranges: str | None = None
) -> list[tuple[Path, int, str]]:
    """
    Split a PDF into one or more files according to page_ranges.
    Returns a list of (output_path, page_count, suggested_filename).
    """
    reader = pypdf.PdfReader(str(input_path))
    total_pages = len(reader.pages)
    if total_pages == 0:
        raise ValueError("PDF has no pages to split.")

    ranges = parse_page_ranges(page_ranges, total_pages)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = Path(base_filename).stem
    results: list[tuple[Path, int, str]] = []

    for idx, page_indices in enumerate(ranges, 1):
        writer = pypdf.PdfWriter()
        for p_idx in page_indices:
            writer.add_page(reader.pages[p_idx])
        if len(page_indices) == 1:
            display_name = f"{stem}_page_{page_indices[0] + 1}.pdf"
        else:
            display_name = f"{stem}_pages_{page_indices[0] + 1}-{page_indices[-1] + 1}.pdf"
        out_file = output_dir / f"{idx}_{display_name}"
        with out_file.open("wb") as f:
            writer.write(f)
        results.append((out_file, len(page_indices), display_name))
    return results


def watermark_pdf(input_path: Path, output_path: Path, watermark_text: str) -> int:
    """Overlay watermark text diagonally across each page of a PDF. Returns page count."""
    text = (watermark_text or "").strip()
    if not text:
        raise ValueError("Watermark text cannot be empty.")
    reader = pypdf.PdfReader(str(input_path))
    total_pages = len(reader.pages)
    if total_pages == 0:
        raise ValueError("PDF has no pages to watermark.")

    writer = pypdf.PdfWriter()
    for page in reader.pages:
        box = page.mediabox
        width = float(box.width)
        height = float(box.height)

        packet = io.BytesIO()
        can = canvas.Canvas(packet, pagesize=(width, height))
        can.setFillColorRGB(0.55, 0.55, 0.55, alpha=0.28)
        font_size = max(24, int(min(width, height) / 10))
        can.setFont("Helvetica-Bold", font_size)
        can.saveState()
        can.translate(width / 2.0, height / 2.0)
        can.rotate(45)
        can.drawCentredString(0, 0, text)
        can.restoreState()
        can.save()

        packet.seek(0)
        watermark_reader = pypdf.PdfReader(packet)
        watermark_page = watermark_reader.pages[0]
        new_page = writer.add_page(page)
        new_page.merge_page(watermark_page)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("wb") as f:
        writer.write(f)
    return total_pages


def convert_image_to_pdf(image_path: Path, output_path: Path) -> int:
    """Convert an image (JPG, PNG) to a 1-page PDF document."""
    with Image.open(str(image_path)) as img:
        rgb_img = img.convert("RGB")
        output_path.parent.mkdir(parents=True, exist_ok=True)
        rgb_img.save(str(output_path), "PDF")
    return 1


def convert_pdf_to_docx(pdf_path: Path, output_path: Path) -> int:
    """Best-effort text extraction from PDF into a new Word (.docx) document."""
    reader = pypdf.PdfReader(str(pdf_path))
    doc = Document()
    total_pages = len(reader.pages)
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if i > 0:
            doc.add_page_break()
        lines = text.splitlines()
        if lines:
            for line in lines:
                if line.strip():
                    doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return total_pages


def convert_docx_to_pdf(docx_path: Path, output_path: Path) -> int:
    """Convert a Word (.docx) document into a PDF document."""
    doc = Document(str(docx_path))
    styles = getSampleStyleSheet()
    story = []
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            escaped = saxutils.escape(text)
            story.append(Paragraph(escaped, styles["Normal"]))
            story.append(Spacer(1, 6))
        else:
            story.append(Spacer(1, 10))
    if not story:
        story.append(Paragraph("(Empty document)", styles["Normal"]))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_doc = SimpleDocTemplate(str(output_path), pagesize=letter)
    pdf_doc.build(story)
    return len(pypdf.PdfReader(str(output_path)).pages)


def extract_text_from_file(file_path: Path | str) -> str:
    """Best-effort text extraction from documents (PDF, DOCX, TXT, etc.) for search indexing."""
    path = Path(file_path)
    if not path.is_file():
        return ""
    suffix = path.suffix.lower()
    # Image files cannot have text extracted without OCR
    if suffix in (".jpg", ".jpeg", ".png"):
        return ""
    if suffix == ".pdf":
        try:
            reader = pypdf.PdfReader(str(path))
            pages = [page.extract_text() or "" for page in reader.pages]
            return "\n".join(p for p in pages if p.strip())
        except Exception:
            return ""
    if suffix == ".docx":
        try:
            doc = Document(str(path))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return ""
    if suffix in (".txt", ".csv", ".rtf", ".odt", ".ods", ".odp", ".doc", ".xls", ".xlsx", ".ppt", ".pptx", ".md", ".json"):
        try:
            return path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return ""
    return ""
